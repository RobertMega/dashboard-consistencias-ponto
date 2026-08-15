from __future__ import annotations

import io
import json
import os
from http.server import BaseHTTPRequestHandler

from openpyxl import Workbook


def _response(handler: BaseHTTPRequestHandler, status: int, body: bytes, content_type: str) -> None:
    handler.send_response(status)
    handler.send_header("Content-Type", content_type)
    handler.send_header("Cache-Control", "private, no-store")
    handler.send_header("Content-Length", str(len(body)))
    handler.end_headers()
    handler.wfile.write(body)


def _workbook(payload: dict) -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.active)
    summary = payload.get("summary", {})
    overview = workbook.create_sheet("Resumo Executivo")
    overview.append(["Indicador", "Valor"])
    for key, value in [("Colaboradores", summary.get("colaboradores", 0)), ("Dias com jornada", summary.get("dias_com_jornada", 0)), ("Pontos a justificar", summary.get("ponto_a_justificar", 0)), ("Marcações ímpares", summary.get("marcacao_impar", 0)), ("Taxa de conformidade", summary.get("taxa_conformidade", 0) / 100), ("Colaboradores críticos", summary.get("colaboradores_criticos", 0)), ("Registros OK", summary.get("registros_ok", 0))]: overview.append([key, value])
    for name, key in [("Ranking Crítico", "maiores_inconsistencias"), ("Ranking Conformidade", "melhores_conformidades")]:
        sheet = workbook.create_sheet(name)
        rows = payload.get("rankings", {}).get(key, [])
        if rows:
            headers = list(rows[0].keys())
            sheet.append(headers)
            for row in rows: sheet.append([row.get(header) for header in headers])
    issues = workbook.create_sheet("Inconsistências")
    issues.append(["Colaborador", "Data", "Tipo", "Severidade", "Status", "Saldo", "Motivo"])
    details = workbook.create_sheet("Detalhamento")
    details.append(["Colaborador", "Data", "Jornada prevista", "Marcações", "Horas totais", "Saldo", "Tipo", "Severidade", "Status", "Motivo"])
    for record in payload.get("records", []):
        types = ", ".join(record.get("inconsistency_types", []))
        details.append([record.get("employee_name"), record.get("work_date"), record.get("planned_minutes"), record.get("marking_count"), record.get("total_minutes"), record.get("original_balance_minutes"), types, record.get("severity"), record.get("status"), record.get("observation") or ""])
        if types not in {"OK", "SEM_JORNADA"}:
            issues.append([record.get("employee_name"), record.get("work_date"), types, record.get("severity"), record.get("status"), record.get("original_balance_minutes"), record.get("observation") or ""])
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]: cell.font = __import__("openpyxl").styles.Font(bold=True, color="FFFFFF"); cell.fill = __import__("openpyxl").styles.PatternFill("solid", fgColor="B91C1C")
    output = io.BytesIO(); workbook.save(output); return output.getvalue()


class handler(BaseHTTPRequestHandler):
    def do_POST(self) -> None:  # noqa: N802
        if self.headers.get("x-parser-token") != os.environ.get("PARSER_SERVICE_TOKEN"):
            _response(self, 401, b"Unauthorized", "text/plain")
            return
        length = int(self.headers.get("content-length", "0"))
        try:
            body = json.loads(self.rfile.read(length))
            _response(self, 200, _workbook(body), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        except Exception:
            _response(self, 422, b"Export failed", "text/plain")

    def do_GET(self) -> None:  # noqa: N802
        _response(self, 405, b"Method not allowed", "text/plain")


__all__ = ["handler"]
