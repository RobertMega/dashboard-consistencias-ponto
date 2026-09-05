from __future__ import annotations

from pathlib import Path
from typing import Any


def _read_xlsx(path: Path) -> list[list[tuple[Any, ...]]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover
        raise RuntimeError("Instale as dependências com: pip install -r automation/requirements.txt") from exc

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        return [
            [tuple(row) for row in sheet.iter_rows(max_col=max(sheet.max_column, 8), values_only=True)]
            for sheet in workbook.worksheets
        ]
    finally:
        workbook.close()


def _read_xls(path: Path) -> list[list[tuple[Any, ...]]]:
    try:
        import xlrd
    except ImportError as exc:  # pragma: no cover
        raise RuntimeError("Instale xlrd com: pip install -r automation/requirements.txt") from exc

    workbook = xlrd.open_workbook(path, on_demand=True)
    try:
        sheets: list[list[tuple[Any, ...]]] = []
        for sheet in workbook.sheets():
            rows: list[tuple[Any, ...]] = []
            for row_index in range(sheet.nrows):
                values: list[Any] = []
                for cell in sheet.row(row_index):
                    if cell.ctype == xlrd.XL_CELL_DATE:
                        values.append(xlrd.xldate_as_datetime(cell.value, workbook.datemode))
                    elif cell.ctype == xlrd.XL_CELL_EMPTY:
                        values.append(None)
                    else:
                        values.append(cell.value)
                rows.append(tuple(values))
            sheets.append(rows)
        return sheets
    finally:
        workbook.release_resources()


def read_workbook_sheets(path: str | Path) -> list[list[tuple[Any, ...]]]:
    source = Path(path)
    if source.suffix.lower() == ".xlsx":
        return _read_xlsx(source)
    if source.suffix.lower() == ".xls":
        return _read_xls(source)
    raise ValueError("Formato Excel não suportado.")
