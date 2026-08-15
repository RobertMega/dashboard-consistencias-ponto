from __future__ import annotations

from datetime import datetime
from pathlib import Path

from automation.parser.metrics import calculate_metrics, rankings
from automation.parser.models import NormalizedRecord, minutes_to_hhmm


HEADER_COLOR = "B91C1C"
SUBHEADER_COLOR = "FEE2E2"


def _detail_rows(records: list[NormalizedRecord]) -> list[list[object]]:
    headers = [
        "Colaborador", "Data", "Dia", "1ª Entrada", "1ª Saída", "2ª Entrada", "2ª Saída",
        "3ª Entrada", "3ª Saída", "Marcações esperadas", "Marcações realizadas", "Horas totais",
        "Horas previstas", "Adicional noturno", "Saldo original", "Saldo diário calculado",
        "Motivo/Observação", "Tipo de inconsistência", "Severidade", "Status", "Página origem",
    ]
    rows = [headers]
    for record in records:
        rows.append([
            record.employee_name,
            record.work_date,
            record.weekday,
            record.markings.get("entrada_1") or "",
            record.markings.get("saida_1") or "",
            record.markings.get("entrada_2") or "",
            record.markings.get("saida_2") or "",
            record.markings.get("entrada_3") or "",
            record.markings.get("saida_3") or "",
            record.expected_marking_count if record.expected_marking_count is not None else "",
            record.marking_count,
            minutes_to_hhmm(record.total_minutes),
            minutes_to_hhmm(record.planned_minutes),
            minutes_to_hhmm(record.additional_night_minutes),
            minutes_to_hhmm(record.original_balance_minutes),
            minutes_to_hhmm(record.calculated_daily_balance_minutes),
            record.observation or "",
            ", ".join(record.inconsistency_types),
            record.severity,
            record.status,
            record.source_page or "",
        ])
    return rows


def _inconsistency_rows(records: list[NormalizedRecord]) -> list[list[object]]:
    rows = [["Data", "Colaborador", "Tipo", "Severidade", "Horas previstas", "Marcações realizadas", "Motivo", "Página"]]
    for record in records:
        for issue in record.inconsistency_types:
            if issue in {"OK", "SEM_JORNADA"}:
                continue
            rows.append([
                record.work_date,
                record.employee_name,
                issue,
                record.severity,
                minutes_to_hhmm(record.planned_minutes),
                record.marking_count,
                record.observation or "",
                record.source_page or "",
            ])
    return rows


def _write_rows(sheet, rows: list[list[object]], header_row: int = 1) -> None:
    for row in rows:
        sheet.append(row)
    for cell in sheet[header_row]:
        cell.font = __import__("openpyxl").styles.Font(color="FFFFFF", bold=True)
        cell.fill = __import__("openpyxl").styles.PatternFill("solid", fgColor=HEADER_COLOR)
        cell.alignment = __import__("openpyxl").styles.Alignment(horizontal="center")
    sheet.freeze_panes = f"A{header_row + 1}"
    sheet.auto_filter.ref = sheet.dimensions
    for column in sheet.columns:
        letter = column[0].column_letter
        max_length = min(max(len(str(cell.value or "")) for cell in column) + 2, 42)
        sheet.column_dimensions[letter].width = max(max_length, 12)


def _format_dates(sheet) -> None:
    for row in sheet.iter_rows():
        for cell in row:
            if hasattr(cell.value, "year"):
                cell.number_format = "dd/mm/yyyy"


def export_excel(records: list[NormalizedRecord], output_path: str | Path, source_pdf: str | None = None) -> Path:
    try:
        from openpyxl import Workbook
    except ImportError as exc:  # pragma: no cover - environment dependent
        raise RuntimeError("Instale as dependências com: pip install -r automation/requirements.txt") from exc

    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    metrics = calculate_metrics(records)
    ranking_data = rankings(records)
    workbook = Workbook()
    workbook.remove(workbook.active)

    dashboard = workbook.create_sheet("Dashboard")
    dashboard.append(["RELATÓRIO GERENCIAL — CONSISTÊNCIAS DE PONTO"])
    dashboard.append(["Indicador", "Valor"])
    indicators = [
        ("Colaboradores", metrics["colaboradores"]),
        ("Dias com jornada", metrics["dias_com_jornada"]),
        ("Pontos a justificar", metrics["ponto_a_justificar"]),
        ("Marcações ímpares", metrics["marcacao_impar"]),
        ("Marcações incompletas", metrics["marcacoes_incompletas"]),
        ("Mais que o previsto", metrics["mais_que_previsto"]),
        ("Ponto sem jornada", metrics["ponto_sem_jornada"]),
        ("Registros OK", metrics["registros_ok"]),
        ("Colaboradores críticos", metrics["colaboradores_criticos"]),
        ("Saldos finais negativos", metrics["saldos_finais_negativos"]),
        ("Taxa de conformidade", metrics["taxa_conformidade"] / 100),
    ]
    for row in indicators:
        dashboard.append(list(row))
    dashboard.append([])
    dashboard.append(["Distribuição por Horas previstas", "Registros"])
    for hours, count in metrics["distribuicao_horas_previstas"].items():
        dashboard.append([hours, count])
    dashboard.append([])
    dashboard.append(["Inconsistências por dia", "Total", "Pontos a justificar", "Marcação ímpar", "Marcações incompletas", "Mais que o previsto"])
    for day, row in metrics["inconsistencias_por_dia"].items():
        dashboard.append([day, row["total"], row["ponto_a_justificar"], row["marcacao_impar"], row["marcacoes_incompletas"], row["mais_que_previsto"]])
    _write_rows(dashboard, [], 2)
    for cell in dashboard[1]:
        cell.font = __import__("openpyxl").styles.Font(color="FFFFFF", bold=True)
        cell.fill = __import__("openpyxl").styles.PatternFill("solid", fgColor=HEADER_COLOR)
    dashboard["A1"].font = __import__("openpyxl").styles.Font(bold=True, size=16, color=HEADER_COLOR)
    dashboard.column_dimensions["A"].width = 34
    dashboard.column_dimensions["B"].width = 18
    _format_dates(dashboard)

    detailed = workbook.create_sheet("Base Detalhada")
    _write_rows(detailed, _detail_rows(records))
    _format_dates(detailed)

    issues = workbook.create_sheet("Inconsistências")
    _write_rows(issues, _inconsistency_rows(records))
    _format_dates(issues)

    employees = workbook.create_sheet("Resumo Colaborador")
    employee_headers = ["Colaborador", "Inconsistências", "Críticas", "Atenção", "Registros OK", "Dias analisados", "Conformidade", "Status"]
    employee_rows = [employee_headers] + [[row["employee_name"], row["inconsistencias"], row["criticas"], row["atencao"], row["registros_ok"], row["dias_analisados"], row["conformidade"] / 100, row["status"]] for row in metrics["saldo_finais"]]
    _write_rows(employees, employee_rows)

    ranking_sheet = workbook.create_sheet("Ranking")
    ranking_headers = ["Ranking", "Posição", "Colaborador", "Inconsistências", "Críticas", "Atenção", "Conformidade", "Registros OK", "Dias analisados", "Status"]
    ranking_rows = [ranking_headers]
    for ranking_name, label in (("maiores_inconsistencias", "Maiores inconsistências"), ("melhores_conformidades", "Melhores conformidades")):
        ranking_rows.extend([[label, row["posicao"], row["employee_name"], row["inconsistencias"], row["criticas"], row["atencao"], row["conformidade"] / 100, row["registros_ok"], row["dias_analisados"], row["status"]] for row in ranking_data[ranking_name]])
    _write_rows(ranking_sheet, ranking_rows)

    balances = workbook.create_sheet("Saldos Finais")
    balance_rows = [["Colaborador", "Saldo final", "Data saldo final", "Crítico"]] + [[row["employee_name"], minutes_to_hhmm(row["saldo_final"]), row["data_saldo_final"] or "", "Sim" if row["critico"] else "Não"] for row in metrics["saldo_finais"]]
    _write_rows(balances, balance_rows)

    history = workbook.create_sheet("Histórico Atualizações")
    history_rows = [["Atualização", "Origem", "Período", "Páginas", "Registros", "Observação"], [datetime.now(), source_pdf or "", "", "", metrics["registros"], "Gerado pelo parser e regras oficiais"]]
    _write_rows(history, history_rows)
    _format_dates(history)

    dictionary = workbook.create_sheet("Dicionário")
    dictionary_rows = [
        ["Campo", "Descrição"],
        ["Horas previstas", "Carga planejada agregada existente no PDF."],
        ["Marcações realizadas", "Valores das seis colunas de Pontos; vazios não contam."],
        ["Marcações esperadas", "2 até 04:00; 4 acima de 04:00; regra centralizada no normalizador."],
        ["Inconsistências", "Classificações produzidas pelas regras de negócio."],
        ["Saldo final", "Saldo do registro cronologicamente mais recente disponível por colaborador."],
        ["Conformidade", "Registros OK dividido pelos dias com jornada."],
    ]
    _write_rows(dictionary, dictionary_rows)

    for sheet in workbook.worksheets:
        for cell in sheet[1]:
            if cell.value is not None and sheet.title != "Dashboard":
                cell.fill = __import__("openpyxl").styles.PatternFill("solid", fgColor=HEADER_COLOR)
                cell.font = __import__("openpyxl").styles.Font(color="FFFFFF", bold=True)

    workbook.save(output)
    return output
